"""Order numbers and listing ids for eBay purchases, from eBay's purchase history.

The owner, 2026-09-25: 1,234 eBay purchases were recorded without an order
number, so they could not be found or checked against eBay. Their purchase
history -- workbooks from the "eBay Download History" Chrome extension, one per
year, a row per line bought (OrderNumber, OrderDate, ItemID, Seller, ItemName,
ItemPrice, ...) -- supplies both the order number and eBay's item id.

**The item id is the key, not the words.** Every one of those purchases links
its eBay listing (`https://www.ebay.com/itm/<item id>`), and so do most items
(`listing_url`); the history has the item id on every line. Measured against
the 1,738 eBay purchases that already had a number, the item id gives back the
same order number for 1,680; of the 58 that disagree, 8 are typos in the
stored number and the rest name another order of the same day -- so those are
listed for a person, never changed here.

What the pass does, all in one transaction:

1. **`sellers_item_id`** on every item whose own link, or failing that its
   purchase's, names an eBay listing. Only where it is empty.
2. **Order numbers.** An eBay purchase with none takes the one order its items'
   ids appear in (a line of that date is preferred when an id was bought
   twice). No order, or more than one, is listed for a person.
3. **One purchase per eBay order.** An order of several listings was recorded
   as a purchase per listing; the database allows one purchase per vendor's
   order number, and an eBay order is one purchase. So the purchases an order
   was split into are merged: their items move to one (the purchase already
   holding the number, else the oldest), the emptied ones are deleted, and each
   item keeps its own listing's id. A merged purchase's link becomes the
   order's page on eBay.

Every item whose order number or listing id changes gets a row in its History
(`item_field_change`) under the person named by `--by`. Dry run by default;
`--commit` writes. `--review` writes a workbook of what was done and of
everything left for a person.

    python -m app.ebay_orders FILE... [--commit --by EMAIL] [--review OUT.xlsx]
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .database import SessionLocal
from .models import InventoryItem, ItemFieldChange, PurchaseOrder, User, Vendor

__all__ = [
    "Line",
    "Plan",
    "apply",
    "item_id_of",
    "plan",
    "read_history",
]

#: The vendor eBay purchases are recorded under.
EBAY = "ebay.com"
#: eBay's listing link: `/itm/<id>`, or `/itm/<title>/<id>` in older links.
_ITEM_ID = re.compile(r"ebay\.[a-z.]+/itm/(?:[^/?#]*/)?(\d{9,15})", re.IGNORECASE)
#: An order's page on eBay, as a merged purchase links it.
ORDER_PAGE = "https://order.ebay.com/ord/show?orderId={}"
#: The history's columns this pass reads.
_COLUMNS = ("OrderNumber", "OrderDate", "ItemID", "Seller", "ItemName", "ItemPrice")


@dataclass(frozen=True)
class Line:
    """One line of eBay's purchase history: an item bought on an order."""

    order_number: str
    ordered_on: date
    item_id: str
    seller: str
    name: str
    price: Decimal | None


def item_id_of(url: str | None) -> str | None:
    """The eBay item id in a listing link, or None for any other link."""
    match = _ITEM_ID.search(url or "")
    return match.group(1) if match else None


def _price(value: object) -> Decimal | None:
    try:
        return Decimal(str(value).replace(",", "")) if value not in (None, "") else None
    except InvalidOperation:
        return None


def read_history(paths: Iterable[Path]) -> list[Line]:
    """Every line of the purchase-history workbooks, first sheet of each."""
    lines: list[Line] = []
    for path in paths:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = book.worksheets[0].iter_rows(values_only=True)
            header = [str(h) if h is not None else "" for h in next(rows, ())]
            missing = [c for c in _COLUMNS if c not in header]
            if missing:
                raise ValueError(f"{path.name}: no column {', '.join(missing)}")
            at = {name: header.index(name) for name in _COLUMNS}
            for values in rows:
                if not values or not values[at["OrderNumber"]]:
                    continue
                lines.append(
                    Line(
                        order_number=str(values[at["OrderNumber"]]).strip(),
                        ordered_on=datetime.strptime(
                            str(values[at["OrderDate"]]).strip(), "%b %d, %Y"
                        ).date(),
                        item_id=str(values[at["ItemID"]]).strip(),
                        seller=str(values[at["Seller"]] or ""),
                        name=str(values[at["ItemName"]] or ""),
                        price=_price(values[at["ItemPrice"]]),
                    )
                )
        finally:
            book.close()
    return lines


@dataclass
class Plan:
    """What the pass would do, and what it leaves for a person."""

    #: item id -> listing id to set.
    listing_ids: dict[int, str] = field(default_factory=dict)
    #: order number -> (surviving purchase id, purchases merged into it).
    orders: dict[str, tuple[int, list[int]]] = field(default_factory=dict)
    #: order number -> its date in the history.
    order_dates: dict[str, date] = field(default_factory=dict)
    #: (purchase id, why) for purchases left without a number.
    unmatched: list[tuple[int, str]] = field(default_factory=list)
    #: (purchase id, stored number, the history's number) where they differ.
    disagreements: list[tuple[int, str, str]] = field(default_factory=list)


def _orders_for(
    ids: set[str], on: date | None, by_item: dict[str, list[Line]]
) -> set[str]:
    """The orders these listing ids were bought on; one of `on`'s date first."""
    lines = [line for i in ids for line in by_item.get(i, [])]
    dated = {line.order_number for line in lines if line.ordered_on == on}
    return dated or {line.order_number for line in lines}


def plan(db: Session, lines: Sequence[Line]) -> Plan:
    """Decide every change from the history, writing nothing."""
    by_item: dict[str, list[Line]] = defaultdict(list)
    for line in lines:
        by_item[line.item_id].append(line)
    out = Plan()
    for line in lines:
        out.order_dates.setdefault(line.order_number, line.ordered_on)

    vendor_id = db.scalar(select(Vendor.id).where(Vendor.name == EBAY))
    if vendor_id is None:
        return out
    purchases = db.scalars(
        select(PurchaseOrder).where(PurchaseOrder.vendor_id == vendor_id)
    ).all()
    items_of: dict[int, list[InventoryItem]] = defaultdict(list)
    for item in db.scalars(
        select(InventoryItem).where(
            InventoryItem.purchase_order_id.in_([p.id for p in purchases])
        )
    ):
        if item.purchase_order_id is not None:
            items_of[item.purchase_order_id].append(item)

    holder = {p.order_number: p.id for p in purchases if p.order_number}
    groups: dict[str, list[int]] = defaultdict(list)
    for purchase in sorted(purchases, key=lambda p: p.id):
        own = item_id_of(purchase.source_url)
        ids: set[str] = set()
        for item in items_of[purchase.id]:
            listing = item.sellers_item_id or item_id_of(item.listing_url) or own
            if listing:
                ids.add(listing)
                if item.sellers_item_id is None:
                    out.listing_ids[item.id] = listing
        if own:
            ids.add(own)
        found = _orders_for(ids, purchase.ordered_on, by_item) if ids else set()
        if purchase.order_number:
            if len(found) == 1 and purchase.order_number not in found:
                out.disagreements.append(
                    (purchase.id, purchase.order_number, next(iter(found)))
                )
            continue
        if not ids:
            out.unmatched.append((purchase.id, "no eBay listing link"))
        elif not found:
            out.unmatched.append((purchase.id, "listing not in the purchase history"))
        elif len(found) > 1:
            out.unmatched.append(
                (purchase.id, "in several orders: " + ", ".join(sorted(found)))
            )
        else:
            groups[next(iter(found))].append(purchase.id)

    for number, members in groups.items():
        survivor = holder.get(number, min(members))
        out.orders[number] = (survivor, sorted(m for m in members if m != survivor))
    return out


def apply(db: Session, todo: Plan, user_id: int) -> dict[str, int]:
    """Make the planned changes and log each item's; the caller commits."""
    now = datetime.now(UTC)
    logged: list[ItemFieldChange] = []

    def log(item_id: int, field_name: str, old: object, new: object) -> None:
        logged.append(
            ItemFieldChange(
                inventory_item_id=item_id,
                field_name=field_name,
                old_value=old,
                new_value=new,
                changed_by_id=user_id,
                changed_at=now,
            )
        )

    for item_id, listing in todo.listing_ids.items():
        item = db.get(InventoryItem, item_id)
        assert item is not None
        item.sellers_item_id = listing
        log(item_id, "sellers_item_id", None, listing)

    moved = deleted = numbered = 0
    for number, (survivor_id, merged) in todo.orders.items():
        survivor = db.get(PurchaseOrder, survivor_id)
        assert survivor is not None
        members = [survivor_id, *merged]
        for item in db.scalars(
            select(InventoryItem).where(InventoryItem.purchase_order_id.in_(members))
        ):
            old = db.get(PurchaseOrder, item.purchase_order_id)
            assert old is not None
            if old.order_number != number:
                log(item.id, "order_number", old.order_number, number)
            if item.purchase_order_id != survivor_id:
                item.purchase_order_id = survivor_id
                moved += 1
        if survivor.order_number != number:
            survivor.order_number = number
            numbered += 1
        if survivor.ordered_on is None:
            survivor.ordered_on = todo.order_dates.get(number)
        if merged:
            survivor.source_url = ORDER_PAGE.format(number)
            db.flush()  # the items point at the survivor before the rest go
            for purchase_id in merged:
                gone = db.get(PurchaseOrder, purchase_id)
                assert gone is not None
                # Expired first: a loaded `items` list would still hold the
                # moved items, and deleting the purchase would then null their
                # purchase -- SQLAlchemy's default for a parent's children.
                db.expire(gone)
                left = db.scalar(
                    select(InventoryItem.id)
                    .where(InventoryItem.purchase_order_id == purchase_id)
                    .limit(1)
                )
                if left is not None:
                    raise RuntimeError(
                        f"purchase {purchase_id} still holds item {left}"
                    )
                db.delete(gone)
                deleted += 1
    db.add_all(logged)
    db.flush()
    return {
        "listing ids set": len(todo.listing_ids),
        "purchases numbered": numbered,
        "items moved": moved,
        "purchases merged away": deleted,
        "history rows": len(logged),
    }


def write_review(db: Session, todo: Plan, lines: Sequence[Line], path: Path) -> None:
    """A workbook of what was done and of what is left for a person."""
    by_order: dict[str, list[Line]] = defaultdict(list)
    for line in lines:
        by_order[line.order_number].append(line)

    def purchase_row(purchase_id: int) -> list[object]:
        purchase = db.get(PurchaseOrder, purchase_id)
        items = db.scalars(
            select(InventoryItem)
            .where(InventoryItem.purchase_order_id == purchase_id)
            .order_by(InventoryItem.id)
        ).all()
        return [
            purchase_id,
            purchase.ordered_on if purchase else None,
            purchase.source_url if purchase else None,
            ", ".join(i.item_code for i in items),
            " | ".join(i.source_title for i in items)[:300],
        ]

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Numbered"
    sheet.append(["order number", "purchase", "merged from", "eBay date", "eBay items"])
    for number, (survivor, merged) in sorted(todo.orders.items()):
        sheet.append(
            [
                number,
                survivor,
                ", ".join(map(str, merged)),
                todo.order_dates.get(number),
                " | ".join(line.name for line in by_order[number])[:300],
            ]
        )
    left = book.create_sheet("Needs you")
    left.append(["purchase", "date", "link", "items", "titles", "why"])
    for purchase_id, why in todo.unmatched:
        left.append([*purchase_row(purchase_id), why])
    differ = book.create_sheet("Numbers that disagree")
    differ.append(
        [
            "purchase",
            "date",
            "link",
            "items",
            "titles",
            "stored number",
            "eBay's number for its listing",
            "eBay's items on that order",
            "the stored number on eBay is",
        ]
    )
    for purchase_id, stored, found in todo.disagreements:
        elsewhere = by_order.get(stored)
        differ.append(
            [
                *purchase_row(purchase_id),
                stored,
                found,
                " | ".join(line.name for line in by_order[found])[:300],
                " | ".join(line.name for line in elsewhere)[:300]
                if elsewhere
                else "not in the history (a typo?)",
            ]
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def main(argv: Sequence[str] | None = None) -> int:
    """Report, or with --commit apply, the order numbers and listing ids."""
    parser = argparse.ArgumentParser(prog="ebay_orders", description=__doc__)
    parser.add_argument(
        "files", nargs="+", type=Path, help="purchase-history workbooks"
    )
    parser.add_argument("--commit", action="store_true", help="write the changes")
    parser.add_argument("--by", help="the person the History rows name (email)")
    parser.add_argument("--review", type=Path, help="write a review workbook here")
    args = parser.parse_args(argv)
    if args.commit and not args.by:
        parser.error("--commit needs --by: every change is logged under a person")

    lines = read_history(args.files)
    with SessionLocal() as db:
        todo = plan(db, lines)
        merged = sum(len(m) for _, m in todo.orders.values())
        orders = len({x.order_number for x in lines})
        print(f"history: {len(lines)} lines, {orders} orders")
        print(f"listing ids to set: {len(todo.listing_ids)}")
        print(f"orders to number: {len(todo.orders)}, merging {merged} purchases")
        print(f"left for a person: {len(todo.unmatched)}")
        print(f"stored numbers that disagree: {len(todo.disagreements)}")
        if args.review:
            write_review(db, todo, lines, args.review)
            print(f"review workbook: {args.review}")
        if not args.commit:
            print("dry run: nothing written (--commit --by EMAIL to apply)")
            return 0
        user_id = db.scalar(select(User.id).where(User.email == args.by))
        if user_id is None:
            print(f"refused: no account {args.by}", file=sys.stderr)
            return 2
        for name, count in apply(db, todo, user_id).items():
            print(f"  {name}: {count}")
        db.commit()
    return 0


if __name__ == "__main__":
    sys.exit(main())
