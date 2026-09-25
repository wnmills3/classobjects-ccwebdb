"""The whole database as one Excel workbook, and back again.

A backup a person can open, read and correct (owner, 2026-09-24): the
database is the record. One sheet per
table, every stored column, the ids and foreign keys exactly as stored, so an
import rebuilds the same database -- relationships and all.

    python -m app.workbook_backup export [--out FILE]      live -> workbook
    python -m app.workbook_backup import FILE --to URL     workbook -> database
        [--unknown-for-missing]                            broken link -> Unknown row
    python -m app.workbook_backup compare URL              live vs URL, row by row
    python -m app.workbook_backup widths FILE              remember FILE's column widths

**The tables come from the database itself**, by reflection, not from the
models: a table that has no model would be lost silently by a model-driven
backup.

**What a cell holds.** Numbers and dates as Excel numbers and dates;
timestamps as ISO text, so their time zone survives; JSON as JSON text;
enums as their value; a NULL as an empty cell and an empty string as `""`
(Excel cannot tell the two apart otherwise). A decimal is rounded back to
its column's scale on import. Computed columns (`total_cost`, `sales_tax`,
...) are written for reading, headed "(computed)", and ignored on import --
the database recomputes them.

**Importing** needs an empty database at the same migration revision:
create one and run `alembic upgrade head` against it. Any rows the
migrations put there are removed first, the tables are loaded in
foreign-key order (a column that points at its own table, or at one loaded
later, is filled in a second pass), and every id sequence is moved past the
highest id. It refuses the live database: restore into a new one, compare
it, then switch to it.

**A broken link is refused, all of them at once.** Every foreign key is
checked against the workbook's own rows before anything is written, and the
refusal lists each row that points at nothing. With `--unknown-for-missing`
(owner, 2026-09-24) a link into a *vocabulary* -- a table with a code and a
label, such as `grade` or `mint` -- is pointed at that vocabulary's Unknown
row instead, and every substitution is reported. The Unknown row is the one
coded `unknown` (or labeled Unknown) if the sheet has one; otherwise it is
added, with id 0 -- ids start at 1, so 0 is never a real row's -- and
`source` manual, so a seed load leaves it alone. A vocabulary whose rows need
more than a code and a label (a denomination's face value, a mint's mark) is
not given one by guesswork: add an `unknown` row to its sheet. A link into
anything else -- an item, an order -- is still refused; there is no Unknown
item.

**Column widths are remembered** (owner, 2026-09-24): each export sizes its
columns from `data/workbook_widths.json`, keyed by sheet and column *name*
rather than letter, so a width follows its column when a migration adds or
moves one. Resize columns in an export, save it, and `widths FILE` records
them; a sheet it finds widths on replaces that sheet's entry, the rest stay.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import (
    JSON,
    Boolean,
    Column,
    Date,
    DateTime,
    Enum,
    Integer,
    MetaData,
    Numeric,
    Table,
    Text,
    create_engine,
    select,
    text,
    update,
)
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.engine import Connection, Engine, make_url
from sqlalchemy.exc import DBAPIError
from sqlalchemy.sql.elements import ColumnElement
from sqlalchemy.types import TypeEngine

from .config import REPO_ROOT, settings

__all__ = [
    "EMPTY_STRING",
    "UNKNOWN_CODE",
    "WIDTHS_FILE",
    "Difference",
    "Imported",
    "Substitution",
    "capture_widths",
    "compare",
    "export_workbook",
    "from_cell",
    "import_workbook",
    "load_widths",
    "remember_widths",
    "to_cell",
]

#: How an empty string is written: an empty cell reads back as NULL.
EMPTY_STRING = '""'
#: The header suffix of a computed column, written for reading only.
COMPUTED = " (computed)"
#: Sheets that describe the workbook rather than hold a table.
ABOUT, COLUMNS = "About", "Columns"
#: The remembered column widths: {sheet: {column name: width}}.
WIDTHS_FILE = Path(__file__).resolve().parents[1] / "data" / "workbook_widths.json"
#: Changed when the layout changes, so an old workbook is recognized.
FORMAT = "ccwebdb-workbook-1"
#: The About sheet's first row, which its widths are keyed by.
ABOUT_HEADER = ("format", FORMAT)
#: The migration table: recorded on the About sheet, never loaded.
VERSION_TABLE = "alembic_version"
#: Rows written per insert round trip.
CHUNK = 1000
#: The code of a vocabulary's Unknown row.
UNKNOWN_CODE = "unknown"
#: The columns that make a table a vocabulary (`models.base.ReferenceMixin`).
_VOCABULARY_COLUMNS = frozenset({"code", "label", "sort_order", "is_active", "source"})
#: Broken links named in a refusal; the count says how many more.
_SHOWN = 25


class WorkbookError(RuntimeError):
    """A workbook that cannot be imported as it stands; the message says where."""


# -- one cell -------------------------------------------------------------------


def _is_json(kind: TypeEngine[Any]) -> bool:
    return isinstance(kind, (JSON, JSONB))


def to_cell(value: object, kind: TypeEngine[Any]) -> object:
    """A stored value as a workbook cell holds it."""
    if value is None:
        return None
    if _is_json(kind):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        return EMPTY_STRING if value == "" else value
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, Decimal):
        # At most 12 digits in this schema, well inside a double's 15.
        return float(value)
    return value


def from_cell(value: object, kind: TypeEngine[Any], where: str) -> object:
    """A workbook cell as the column stores it; `where` names it in an error."""
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    try:
        return _convert(value, kind)
    except (ValueError, ArithmeticError, TypeError, json.JSONDecodeError) as exc:
        raise WorkbookError(f"{where}: {value!r} is not a valid {kind}: {exc}") from exc


def _convert(value: object, kind: TypeEngine[Any]) -> object:
    if _is_json(kind):
        # A JSON null is a value, not SQL NULL: sent as JSON's null.
        decoded = json.loads(str(value))
        return JSON.NULL if decoded is None else decoded
    if value == EMPTY_STRING:
        return ""
    if isinstance(kind, Boolean):
        if isinstance(value, bool):
            return value
        word = str(value).strip().lower()
        if word in {"true", "yes", "1"}:
            return True
        if word in {"false", "no", "0"}:
            return False
        raise ValueError("expected TRUE or FALSE")
    if isinstance(kind, Numeric) and not isinstance(kind, Integer):
        number = Decimal(str(value))
        scale = kind.scale
        return number if scale is None else number.quantize(Decimal(1).scaleb(-scale))
    if isinstance(kind, Integer):
        number = Decimal(str(value))
        if number != number.to_integral_value():
            raise ValueError("expected a whole number")
        return int(number)
    if isinstance(kind, DateTime):
        moment = (
            value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
        )
        if kind.timezone and moment.tzinfo is None:
            # Typed in by hand with no zone: read as this machine's local time.
            moment = moment.astimezone()
        return moment
    if isinstance(kind, Date):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return date.fromisoformat(str(value))
    if isinstance(kind, Enum):
        return str(value)
    return value if isinstance(value, str) else str(value)


# -- the schema -----------------------------------------------------------------


def reflect(engine: Engine) -> MetaData:
    """Every table the database holds, in its own terms."""
    meta = MetaData()
    meta.reflect(engine)
    return meta


def _tables(meta: MetaData) -> list[Table]:
    """The data tables in foreign-key order: the migration table is not one."""
    return [t for t in meta.sorted_tables if t.name != VERSION_TABLE]


def _revision(conn: Connection) -> str | None:
    return conn.execute(text(f"SELECT version_num FROM {VERSION_TABLE}")).scalar()


def _stored(table: Table) -> list[Column[Any]]:
    return [c for c in table.columns if c.computed is None]


def _computed(table: Table) -> list[Column[Any]]:
    return [c for c in table.columns if c.computed is not None]


def _order(table: Table) -> list[Column[Any]]:
    return list(table.primary_key.columns) or list(table.columns)


# -- export ---------------------------------------------------------------------


def default_path() -> Path:
    """Beside the other backups, outside the repository, named by the time."""
    folder = REPO_ROOT.parent / "ccwebdb-backups"
    return folder / f"ccwebdb_{datetime.now(UTC).astimezone():%Y%m%d_%H%M%S}.xlsx"


def _as_read(column: Column[Any]) -> ColumnElement[Any]:
    """A JSON column is read as its JSON text, the rest as they are.

    As text, a JSON null reads as `null` while SQL NULL reads as nothing;
    decoded, both would arrive as Python None and the workbook could not
    tell them apart.
    """
    return column.cast(Text()) if _is_json(column.type) else column


def _read_type(column: Column[Any]) -> TypeEngine[Any]:
    """The type of what `_as_read` returns: text for JSON, else the column's."""
    return Text() if _is_json(column.type) else column.type


def _text_cell(sheet: Worksheet, value: object) -> object:
    """A cell that stays text: a value starting with "=" is not a formula."""
    if isinstance(value, str) and value.startswith("="):
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"
        return cell
    return value


Widths = dict[str, dict[str, float]]


def load_widths(path: Path = WIDTHS_FILE) -> Widths:
    """The remembered widths, or none if nothing has been remembered yet."""
    if not path.exists():
        return {}
    loaded: Widths = json.loads(path.read_text(encoding="utf-8"))
    return loaded


def capture_widths(workbook: Path) -> Widths:
    """The widths a person set in `workbook`, by sheet and column-heading name.

    Only widths set by hand (Excel's custom width); a column left at the
    default is not recorded. Excel may save one width for a run of adjacent
    columns, so each run is expanded to every column it covers.
    """
    book = load_workbook(workbook)
    out: Widths = {}
    for sheet in book.worksheets:
        names = {
            cell.column: str(cell.value)
            for cell in next(sheet.iter_rows(min_row=1, max_row=1), ())
            if cell.value is not None
        }
        found: dict[str, float] = {}
        for dimension in sheet.column_dimensions.values():
            if not dimension.customWidth or not dimension.width:
                continue
            for index in range(dimension.min or 0, (dimension.max or 0) + 1):
                if index in names:
                    found[names[index]] = round(float(dimension.width), 2)
        if found:
            out[sheet.title] = dict(sorted(found.items()))
    book.close()
    return out


def remember_widths(workbook: Path, path: Path = WIDTHS_FILE) -> Widths:
    """Merge `workbook`'s widths into the remembered ones; what is stored now.

    A sheet with widths in `workbook` replaces that sheet's entry outright --
    the owner's latest sizing of it -- and sheets it has no widths on keep
    theirs.
    """
    merged = load_widths(path)
    merged.update(capture_widths(workbook))
    ordered = dict(sorted(merged.items()))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(ordered, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    return ordered


def _size(sheet: Worksheet, header: Sequence[object], widths: dict[str, float]) -> None:
    """Size a sheet's columns from the remembered widths, before its first row.

    A write-only sheet writes its column widths ahead of the rows, so this must
    run before the header is appended.
    """
    for index, name in enumerate(header, start=1):
        width = widths.get(str(name))
        if width:
            sheet.column_dimensions[get_column_letter(index)].width = width


def export_workbook(
    engine: Engine, path: Path, widths: Widths | None = None
) -> dict[str, int]:
    """Write every table to `path`; the rows written, per table.

    `widths` sizes the columns; by default the remembered ones.
    """
    sizes = load_widths() if widths is None else widths
    meta = reflect(engine)
    tables = _tables(meta)
    book = Workbook(write_only=True)
    counts: dict[str, int] = {}
    with engine.connect() as conn:
        revision = _revision(conn)
        about = book.create_sheet(ABOUT)
        _size(about, ABOUT_HEADER, sizes.get(ABOUT, {}))
        columns = book.create_sheet(COLUMNS)
        column_header = [
            "table",
            "column",
            "type",
            "nullable",
            "computed",
            "references",
        ]
        _size(columns, column_header, sizes.get(COLUMNS, {}))
        columns.append(column_header)
        for table in tables:
            for column in table.columns:
                refs = ", ".join(
                    f"{fk.column.table.name}.{fk.column.name}"
                    for fk in column.foreign_keys
                )
                columns.append(
                    [
                        table.name,
                        column.name,
                        str(column.type),
                        column.nullable,
                        column.computed is not None,
                        refs,
                    ]
                )
        for table in tables:
            sheet = book.create_sheet(table.name)
            stored, computed = _stored(table), _computed(table)
            header = [c.name for c in stored] + [c.name + COMPUTED for c in computed]
            _size(sheet, header, sizes.get(table.name, {}))
            sheet.append(header)
            n = 0
            everything = [*stored, *computed]
            result = conn.execution_options(yield_per=CHUNK).execute(
                select(*(_as_read(c) for c in everything)).order_by(*_order(table))
            )
            for row in result:
                cells = [
                    _text_cell(sheet, to_cell(value, _read_type(column)))
                    for value, column in zip(row, everything, strict=True)
                ]
                sheet.append(cells)
                n += 1
            counts[table.name] = n
        about.append(list(ABOUT_HEADER))
        about.append(["exported", datetime.now(UTC).astimezone().isoformat()])
        about.append(["database", make_url(str(engine.url)).database])
        about.append(["migration revision", revision])
        about.append(["tables", len(tables)])
        about.append(["empty cell", "NULL -- no value"])
        about.append(["empty string", f"written as {EMPTY_STRING}"])
        about.append(
            ["timestamps", "ISO text with the time zone; typed with none = local time"]
        )
        about.append(
            ["(computed)", "columns written for reading; the import ignores them"]
        )
        about.append([])
        about.append(["table", "rows"])
        for name, n in counts.items():
            about.append([name, n])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return counts


# -- import ---------------------------------------------------------------------


@dataclass(frozen=True)
class _Loaded:
    """One table's rows, converted, and the columns filled in a second pass."""

    table: Table
    rows: list[dict[str, object]]
    deferred: list[str]


def _about(book: Workbook) -> dict[str, object]:
    if ABOUT not in book.sheetnames:
        raise WorkbookError(f"no {ABOUT} sheet: not a ccwebdb workbook")
    out: dict[str, object] = {}
    for row in book[ABOUT].iter_rows(values_only=True):
        if row and row[0] is not None and len(row) > 1:
            out[str(row[0])] = row[1]
            if row[0] == "table":
                break
    if out.get("format") != FORMAT:
        raise WorkbookError(f"format is {out.get('format')!r}, expected {FORMAT!r}")
    return out


def _read_sheet(book: Workbook, table: Table, later: set[str]) -> _Loaded:
    """Convert one sheet to rows for `table`, checking every header."""
    if table.name not in book.sheetnames:
        raise WorkbookError(f"no sheet for table {table.name}")
    rows = book[table.name].iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows, ())]
    stored = {c.name: c for c in _stored(table)}
    kept = [(i, name) for i, name in enumerate(header) if name in stored]
    unknown = [n for n in header if n and n not in stored and not n.endswith(COMPUTED)]
    if unknown:
        raise WorkbookError(f"{table.name}: no such column {', '.join(unknown)}")
    missing = sorted(set(stored) - {name for _, name in kept})
    if missing:
        raise WorkbookError(f"{table.name}: missing column {', '.join(missing)}")
    deferred = sorted(
        {
            fk.parent.name
            for fk in table.foreign_keys
            if fk.column.table is table or fk.column.table.name in later
        }
    )
    out: list[dict[str, object]] = []
    for number, values in enumerate(rows, start=2):
        if values is None or all(v is None for v in values):
            continue
        record: dict[str, object] = {}
        for i, name in kept:
            value = values[i] if i < len(values) else None
            record[name] = from_cell(
                value, stored[name].type, f"{table.name} row {number}, {name}"
            )
        out.append(record)
    return _Loaded(table, out, deferred)


def _refuse_live(url: str) -> None:
    live = make_url(settings.database_url)
    target = make_url(url)
    if (target.host, target.port, target.database) == (
        live.host,
        live.port,
        live.database,
    ):
        raise WorkbookError(
            f"{target.database} is the live database: import into a new one, "
            "compare it, then switch to it"
        )


@dataclass(frozen=True)
class Substitution:
    """A link that pointed at nothing, pointed at its vocabulary's Unknown row."""

    table: str
    row: str
    column: str
    missing: object
    vocabulary: str
    unknown: object


@dataclass(frozen=True)
class Imported:
    """What an import loaded: rows per table, and every Unknown substitution."""

    counts: dict[str, int]
    substituted: list[Substitution]


def import_workbook(
    path: Path, url: str, *, unknown_for_missing: bool = False
) -> Imported:
    """Load the workbook at `path` into the database at `url`.

    A link that points at no row is refused, unless `unknown_for_missing`
    and the link is into a vocabulary: then it points at that vocabulary's
    Unknown row, and the substitution is returned.
    """
    _refuse_live(url)
    book = load_workbook(path, read_only=True, data_only=True)
    engine = create_engine(url)
    try:
        return _load(engine, book, unknown_for_missing=unknown_for_missing)
    finally:
        # Closed whether it loaded or was refused: a pooled connection left
        # open holds the database, and nothing can drop or rename it.
        book.close()
        engine.dispose()


def _is_vocabulary(table: Table) -> bool:
    return {c.name for c in table.columns} >= _VOCABULARY_COLUMNS


def _key(part: _Loaded, row: dict[str, object]) -> str:
    """A row named by its primary key, as a refusal or a report shows it."""
    keys = list(part.table.primary_key.columns) or list(part.table.columns)[:1]
    return ", ".join(f"{k.name} {row.get(k.name)}" for k in keys)


def _unknown_row(part: _Loaded, column: str) -> object:
    """The value of `column` on the vocabulary's Unknown row, added if absent."""
    for row in part.rows:
        code, label = row.get("code"), row.get("label")
        if code == UNKNOWN_CODE or (
            isinstance(label, str) and label.lower() == "unknown"
        ):
            return row[column]
    table = part.table
    primary = list(table.primary_key.columns)
    if [c.name for c in primary] != [column] or not isinstance(
        primary[0].type, Integer
    ):
        raise WorkbookError(
            f"{table.name}: links point at {column}, not its integer id, so no "
            f"Unknown row can be added: add a row coded {UNKNOWN_CODE!r} to its sheet"
        )
    filled = {"code", "label", "sort_order", "is_active", "source", column}
    needed = [
        c.name
        for c in _stored(table)
        if not c.nullable and c.server_default is None and c.name not in filled
    ]
    if needed:
        raise WorkbookError(
            f"{table.name}: an Unknown row needs {', '.join(needed)}, which cannot "
            f"be guessed: add a row coded {UNKNOWN_CODE!r} to its sheet"
        )
    taken = {row.get(column) for row in part.rows}
    new_id = 0
    while new_id in taken:
        new_id -= 1
    orders = [o for r in part.rows if isinstance(o := r.get("sort_order"), int)]
    # Only the columns set here: the database fills the rest, NULL or its own
    # default (`grade.is_plus` is NOT NULL DEFAULT false) -- an explicit None
    # would be refused. `_load` inserts a row of other columns on its own.
    part.rows.append(
        {
            column: new_id,
            "code": UNKNOWN_CODE,
            "label": "Unknown",
            # Last in a sequenced picker, after every real value.
            "sort_order": max(orders, default=0) + 1,
            "is_active": True,
            "source": "manual",
        }
    )
    return new_id


def _check_links(parts: list[_Loaded], unknown_for_missing: bool) -> list[Substitution]:
    """Refuse or substitute every link that points at no row in the workbook.

    Checked here, before anything is written, so a refusal names every broken
    link at once rather than the database's first. A link of several columns
    is left to the database.
    """
    by_name = {part.table.name: part for part in parts}
    present: dict[tuple[str, str], set[object]] = {}

    def values(table: str, column: str) -> set[object]:
        if (table, column) not in present:
            present[table, column] = {r.get(column) for r in by_name[table].rows}
        return present[table, column]

    broken: list[str] = []
    substituted: list[Substitution] = []
    for part in parts:
        for fk in sorted(part.table.foreign_keys, key=lambda f: f.parent.name):
            constraint = fk.constraint
            if (
                constraint is None
                or len(constraint.elements) != 1
                or fk.column.table.name not in by_name
            ):
                continue
            target, column = fk.column.table, fk.column.name
            for row in part.rows:
                value = row.get(fk.parent.name)
                if value is None or value in values(target.name, column):
                    continue
                if unknown_for_missing and _is_vocabulary(target):
                    unknown = _unknown_row(by_name[target.name], column)
                    values(target.name, column).add(unknown)
                    row[fk.parent.name] = unknown
                    substituted.append(
                        Substitution(
                            part.table.name,
                            _key(part, row),
                            fk.parent.name,
                            value,
                            target.name,
                            unknown,
                        )
                    )
                else:
                    broken.append(
                        f"{part.table.name} {_key(part, row)}: {fk.parent.name} "
                        f"{value!r} is no {target.name}.{column}"
                    )
    if broken:
        shown = "; ".join(broken[:_SHOWN])
        more = f" (and {len(broken) - _SHOWN} more)" if len(broken) > _SHOWN else ""
        hint = (
            ""
            if unknown_for_missing
            else (
                " -- with --unknown-for-missing, a link into a vocabulary "
                "points at its Unknown row instead"
            )
        )
        raise WorkbookError(
            f"{len(broken)} link(s) point at nothing: {shown}{more}{hint}"
        )
    return substituted


def _load(engine: Engine, book: Workbook, *, unknown_for_missing: bool) -> Imported:
    """Replace every table's rows with the workbook's, in one transaction."""
    about = _about(book)
    tables = _tables(reflect(engine))
    for table in tables:
        for column in table.columns:
            if _is_json(column.type):
                # None is SQL NULL here; JSON's null comes as JSON.NULL.
                column.type = JSONB(none_as_null=True)
    counts: dict[str, int] = {}
    with engine.begin() as conn:
        revision = _revision(conn)
        if revision != about.get("migration revision"):
            raise WorkbookError(
                f"the workbook is at revision {about.get('migration revision')}, "
                f"the database at {revision}: run `alembic upgrade` to match"
            )
        # Anything the migrations inserted goes: the workbook is the whole record.
        names = ", ".join(
            engine.dialect.identifier_preparer.quote(t.name) for t in tables
        )
        conn.execute(text(f"TRUNCATE {names} RESTART IDENTITY CASCADE"))

        remaining = {t.name for t in tables}
        loaded: list[_Loaded] = []
        for table in tables:
            remaining.discard(table.name)
            loaded.append(_read_sheet(book, table, remaining))
        substituted = _check_links(loaded, unknown_for_missing)
        for part in loaded:
            # One insert per set of columns: a sheet's rows all share one, an
            # added Unknown row names fewer.
            shapes: dict[tuple[str, ...], list[dict[str, object]]] = {}
            for r in part.rows:
                first = {k: (None if k in part.deferred else v) for k, v in r.items()}
                shapes.setdefault(tuple(first), []).append(first)
            for rows in shapes.values():
                for start in range(0, len(rows), CHUNK):
                    _insert(conn, part.table, rows[start : start + CHUNK])
            counts[part.table.name] = len(part.rows)
        for part in loaded:
            _second_pass(conn, part)
        for table in tables:
            _resync(conn, table)
    return Imported(counts, substituted)


def _insert(conn: Connection, table: Table, rows: list[dict[str, object]]) -> None:
    """Insert rows; a refusal names the table and the database's own reason."""
    try:
        conn.execute(table.insert(), rows)
    except DBAPIError as exc:
        reason = str(exc.orig).strip().splitlines()
        raise WorkbookError(
            f"{table.name}: the database refused a row -- {' '.join(reason[:2])}"
        ) from exc


def _second_pass(conn: Connection, part: _Loaded) -> None:
    """Fill the links that could not be written before their targets existed."""
    if not part.deferred:
        return
    keys = list(part.table.primary_key.columns)
    for row in part.rows:
        values = {
            name: row[name] for name in part.deferred if row.get(name) is not None
        }
        if not values:
            continue
        where = [key == row[key.name] for key in keys]
        try:
            conn.execute(update(part.table).where(*where).values(**values))
        except DBAPIError as exc:
            reason = str(exc.orig).strip().splitlines()
            raise WorkbookError(
                f"{part.table.name}: the database refused a link -- "
                f"{' '.join(reason[:2])}"
            ) from exc


def _resync(conn: Connection, table: Table) -> None:
    """Move an id sequence past the highest id loaded, as `app.backup` does."""
    primary = list(table.primary_key.columns)
    if len(primary) != 1 or primary[0].name != "id":
        return
    sequence = conn.execute(
        text("SELECT pg_get_serial_sequence(:t, 'id')"), {"t": table.name}
    ).scalar()
    if sequence is None:
        return
    quoted = conn.dialect.identifier_preparer.quote(table.name)
    conn.execute(
        text(
            "SELECT setval(:s, coalesce((SELECT max(id) FROM " + quoted + "), 1), true)"
        ),
        {"s": sequence},
    )


# -- compare --------------------------------------------------------------------


@dataclass(frozen=True)
class Difference:
    """A table whose contents differ between two databases."""

    table: str
    left_rows: int
    right_rows: int


def _digest(conn: Connection, table: Table) -> tuple[int, str]:
    """Row count and a digest of every row, in key order, as text."""
    quote = conn.dialect.identifier_preparer.quote
    quoted = quote(table.name)
    order = ", ".join(quote(c.name) for c in _order(table))
    # Columns named, in name order: a database built by migrations and one
    # built from the models hold the same columns in different physical
    # orders, and `t::text` would call identical rows different.
    ordered = sorted(table.columns, key=lambda c: c.name)
    fields = ", ".join(f"t.{quote(c.name)}" for c in ordered)
    joined = f"string_agg(ROW({fields})::text, E'\\n' ORDER BY {order})"
    row = conn.execute(
        text(f"SELECT count(*), md5(coalesce({joined}, '')) FROM {quoted} t")
    ).one()
    return int(row[0]), str(row[1])


def compare(left: Engine, right: Engine) -> list[Difference]:
    """Every table whose rows differ, compared cell by cell through a digest."""
    out: list[Difference] = []
    tables = _tables(reflect(left))
    with left.connect() as a, right.connect() as b:
        if _revision(a) != _revision(b):
            out.append(Difference(VERSION_TABLE, 1, 1))
        present = set(reflect(right).tables)
        for table in tables:
            if table.name not in present:
                out.append(Difference(table.name, _digest(a, table)[0], -1))
                continue
            left_digest, right_digest = _digest(a, table), _digest(b, table)
            if left_digest != right_digest:
                out.append(Difference(table.name, left_digest[0], right_digest[0]))
    return out


# -- command line ---------------------------------------------------------------


def _print_counts(counts: dict[str, int]) -> None:
    for name, n in counts.items():
        if n:
            print(f"  {name:<28}{n:>8,}")
    print(f"  {len(counts)} tables, {sum(counts.values()):,} rows")


def _lines(differences: Iterable[Difference]) -> Iterator[str]:
    for d in differences:
        right = "missing" if d.right_rows < 0 else f"{d.right_rows:,}"
        yield f"  DIFFERS {d.table}: {d.left_rows:,} rows vs {right}"


def main(argv: Sequence[str] | None = None) -> int:
    """Export, import or compare."""
    parser = argparse.ArgumentParser(prog="workbook_backup", description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    exp = sub.add_parser("export", help="write the live database to a workbook")
    exp.add_argument("--out", type=Path, help="the file; default ccwebdb-backups/")
    imp = sub.add_parser("import", help="load a workbook into an empty database")
    imp.add_argument("file", type=Path)
    imp.add_argument("--to", required=True, help="the database URL to load into")
    imp.add_argument(
        "--unknown-for-missing",
        action="store_true",
        help="point a vocabulary link that finds no row at its Unknown row",
    )
    cmp_ = sub.add_parser("compare", help="compare the live database with another")
    cmp_.add_argument("url")
    wid = sub.add_parser("widths", help="remember a workbook's column widths")
    wid.add_argument("file", type=Path)
    args = parser.parse_args(argv)

    live = create_engine(settings.database_url)
    try:
        if args.command == "widths":
            stored = remember_widths(args.file)
            columns = sum(len(v) for v in stored.values())
            print(
                f"remembered {columns} column widths on {len(stored)} sheets "
                f"in {WIDTHS_FILE}"
            )
            return 0
        if args.command == "export":
            path = args.out or default_path()
            counts = export_workbook(live, path)
            _print_counts(counts)
            print(f"wrote {path}")
        elif args.command == "import":
            imported = import_workbook(
                args.file, args.to, unknown_for_missing=args.unknown_for_missing
            )
            _print_counts(imported.counts)
            for s in imported.substituted:
                print(
                    f"  UNKNOWN {s.table} {s.row}: {s.column} {s.missing!r} -> "
                    f"{s.vocabulary} Unknown ({s.unknown})"
                )
            if imported.substituted:
                print(f"  {len(imported.substituted)} link(s) set to Unknown")
            print(f"loaded {args.file} into {make_url(args.to).database}")
        else:
            differences = compare(live, create_engine(args.url))
            for line in _lines(differences):
                print(line)
            print(
                "identical"
                if not differences
                else f"{len(differences)} table(s) differ"
            )
            return 1 if differences else 0
    except WorkbookError as exc:
        print(f"refused: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
