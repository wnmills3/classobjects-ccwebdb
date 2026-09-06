"""Reading tabular sources into verbatim rows.

DURABLE. Knows about file formats, never about their meaning.
"""

from __future__ import annotations

import hashlib
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from .profile import RawRow


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _as_text(value: object) -> str | None:
    """Render a cell as text without losing or inventing information.

    Spreadsheet libraries hand back native types; the staging layer stores text
    so that nothing is silently re-typed on the way in. Dates become ISO so they
    round-trip, and floats that are whole numbers keep their integer form rather
    than gaining a spurious '.0'.
    """
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


class XlsxSource:
    """An .xlsx worksheet, read verbatim."""

    kind = "xlsx"

    def __init__(self, path: str | Path, sheet: str | None = None) -> None:
        self.path = Path(path)
        if not self.path.is_file():
            raise FileNotFoundError(self.path)
        self.sheet = sheet
        self._sha256: str | None = None

    @property
    def sha256(self) -> str:
        if self._sha256 is None:
            self._sha256 = sha256_of(self.path)
        return self._sha256

    def read_rows(self, limit: int | None = None) -> Iterator[RawRow]:
        from openpyxl import load_workbook

        wb = load_workbook(self.path, data_only=True, read_only=True)
        try:
            ws = wb[self.sheet] if self.sheet else wb[wb.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
            try:
                header_cells = next(rows)
            except StopIteration:
                return
            headers = [
                (_as_text(h) or f"column_{i}") for i, h in enumerate(header_cells)
            ]

            for offset, raw in enumerate(rows):
                if limit is not None and offset >= limit:
                    break
                # row_number matches what a spreadsheet UI shows: header is 1
                values = {
                    headers[i]: _as_text(raw[i]) if i < len(raw) else None
                    for i in range(len(headers))
                }
                if all(v is None or v == "" for v in values.values()):
                    continue  # skip wholly blank rows
                yield RawRow(row_number=offset + 2, values=values)
        finally:
            wb.close()
