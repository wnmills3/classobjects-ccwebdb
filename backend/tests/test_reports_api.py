"""The reports API: the catalog, a result as JSON, and a result as a workbook."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from app.models import ItemStatus, User
from app.reports import REPORTS
from app.reports.base import Column, ReportResult
from app.reports.purchasing import PR_OUTSTANDING, OutstandingParams
from app.reports.workbook import write_workbook
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from tests.builders import (
    ItemFactory,
    build_bare_item,
    build_purchase_order,
    code_id,
    priced_item,
)
from tests.test_reports_performance import _build_modest_collection
from tests.test_reports_selling import _placed, _venue

_ENDPOINTS = (
    "/api/reports",
    "/api/reports/pr_outstanding",
    "/api/reports/pr_outstanding/workbook",
)


def test_the_catalog_lists_every_report_in_registry_order(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get("/api/reports", headers=admin_headers)
    assert res.status_code == 200
    body = res.json()
    assert [entry["id"] for entry in body] == list(REPORTS)
    for entry in body:
        assert set(entry) == {"id", "group", "title", "purpose", "params"}


def test_cb_holdings_carries_its_choices_and_defaults(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get("/api/reports", headers=admin_headers)
    entry = next(e for e in res.json() if e["id"] == "cb_holdings")
    by_name = {p["name"]: p for p in entry["params"]}
    assert by_name["status"]["type"] == "choice"
    assert by_name["status"]["default"] == "received"
    assert by_name["status"]["label"] == "Status"
    assert "ordered" in by_name["status"]["choices"]
    assert "all" in by_name["status"]["choices"]
    assert by_name["disposition"]["default"] == "held"


def test_pr_outstanding_default_and_label(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get("/api/reports", headers=admin_headers)
    entry = next(e for e in res.json() if e["id"] == "pr_outstanding")
    assert len(entry["params"]) == 1
    param = entry["params"][0]
    assert param["name"] == "overdue_days"
    assert param["type"] == "integer"
    assert param["default"] == 21
    assert param["label"] == "Overdue after (days)"
    assert param["choices"] is None


def test_every_registered_report_runs_and_aligns_rows_and_drills(
    client: TestClient, admin_headers: dict[str, str], db: Session
) -> None:
    """A smoke test over the whole catalog, with just enough data to run."""
    build_bare_item(db)
    order = build_purchase_order(db, vendor_name="Outstanding Vendor", commit=False)
    item = build_bare_item(db, status_id=code_id(db, ItemStatus, "ordered"))
    item.purchase_order_id = order.id
    db.commit()

    for report_id in REPORTS:
        res = client.get(f"/api/reports/{report_id}", headers=admin_headers)
        assert res.status_code == 200, (report_id, res.text)
        body = res.json()
        assert body["id"] == report_id
        assert len(body["drills"]) == len(body["rows"])
        column_keys = {c["key"] for c in body["columns"]}
        assert body["link_column"] in column_keys
        for row in body["rows"]:
            assert set(row) == column_keys


def test_each_reports_links_sit_on_the_column_that_names_what_they_open(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    """A listing links on what it offers, a check on its own name; else column 1."""
    expected = {
        "sl_offered": "offers",
        "dq_issues": "check",
        "dq_completeness": "kind",
        "cb_holdings": "kind",
    }
    for report_id, link_column in expected.items():
        res = client.get(f"/api/reports/{report_id}", headers=admin_headers)
        assert res.status_code == 200, (report_id, res.text)
        assert res.json()["link_column"] == link_column, report_id


def test_a_param_round_trips_into_the_result(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get(
        "/api/reports/pr_outstanding?overdue_days=5", headers=admin_headers
    )
    assert res.status_code == 200
    assert res.json()["params"] == {"overdue_days": 5}


def test_an_unknown_report_is_a_404(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    assert client.get("/api/reports/nope", headers=admin_headers).status_code == 404
    res = client.get("/api/reports/nope", headers=admin_headers)
    assert res.json()["detail"] == "Report not found"
    assert (
        client.get("/api/reports/nope/workbook", headers=admin_headers).status_code
        == 404
    )


def test_a_bad_parameter_value_is_a_422(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get(
        "/api/reports/pr_outstanding?overdue_days=0", headers=admin_headers
    )
    assert res.status_code == 422

    res = client.get("/api/reports/cb_holdings?status=bogus", headers=admin_headers)
    assert res.status_code == 422


def test_an_unknown_query_key_is_a_422(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get("/api/reports/pr_outstanding?nope=1", headers=admin_headers)
    assert res.status_code == 422


def test_anonymous_is_refused_on_every_endpoint(client: TestClient) -> None:
    for path in _ENDPOINTS:
        assert client.get(path).status_code == 401


def test_a_customer_is_refused_on_every_endpoint(
    client: TestClient, customer_headers: dict[str, str]
) -> None:
    for path in _ENDPOINTS:
        assert client.get(path, headers=customer_headers).status_code == 403


def _header_row(sheet: Worksheet, first_label: str) -> int:
    """The row number whose first cell is `first_label` -- the header row."""
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=1):
        if row[0].value == first_label:
            return row[0].row
    raise AssertionError(f"no row starts with {first_label!r}")


def test_the_workbook_rows_equal_the_jsons_rows(
    client: TestClient, admin_headers: dict[str, str], db: Session
) -> None:
    order = build_purchase_order(
        db, vendor_name="Workbook Vendor", order_number="WB-1", commit=False
    )
    item = build_bare_item(
        db,
        status_id=code_id(db, ItemStatus, "ordered"),
        item_cost=Decimal("42.50"),
        shipping_cost=Decimal("0.00"),
        tax_rate=Decimal("0"),
    )
    item.purchase_order_id = order.id
    db.commit()

    json_res = client.get("/api/reports/pr_outstanding", headers=admin_headers)
    assert json_res.status_code == 200
    json_body = json_res.json()
    columns = json_body["columns"]

    wb_res = client.get("/api/reports/pr_outstanding/workbook", headers=admin_headers)
    assert wb_res.status_code == 200
    assert wb_res.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = wb_res.headers["content-disposition"]
    assert disposition.startswith("attachment; filename=")
    assert "pr_outstanding_" in disposition
    assert disposition.endswith('.xlsx"')

    book = load_workbook(BytesIO(wb_res.content))
    assert book.sheetnames == ["Not yet arrived"]
    sheet = book["Not yet arrived"]

    header_row = _header_row(sheet, columns[0]["label"])
    for offset, json_row in enumerate(json_body["rows"], start=1):
        row_index = header_row + offset
        for col_index, column in enumerate(columns, start=1):
            cell = sheet.cell(row=row_index, column=col_index)
            json_value = json_row[column["key"]]
            if column["kind"] == "money":
                assert cell.number_format == "#,##0.00"
                if json_value is None:
                    assert cell.value is None
                else:
                    assert Decimal(str(cell.value)) == Decimal(json_value)
            elif column["kind"] == "date":
                if json_value is None:
                    assert cell.value is None
                else:
                    assert isinstance(cell.value, (date, datetime))
                    assert cell.value.isoformat() == json_value
            else:
                # Excel cannot tell an empty string from a blank cell --
                # the same limitation `workbook_backup` works around with
                # its own sentinel for a round-tripped import; a report
                # export is read-only and a blank cell reads the same to a
                # person either way, so both collapse to `None` here.
                expected = json_value if json_value != "" else None
                assert cell.value == expected

    assert json_body["totals"] is not None
    totals_row = header_row + len(json_body["rows"]) + 1
    money_col_index = next(
        i for i, c in enumerate(columns, start=1) if c["kind"] == "money"
    )
    money_key = columns[money_col_index - 1]["key"]
    totals_cell = sheet.cell(row=totals_row, column=money_col_index)
    assert totals_cell.font.bold is True
    assert totals_cell.number_format == "#,##0.00"
    assert Decimal(str(totals_cell.value)) == Decimal(json_body["totals"][money_key])

    assert json_body["notes"]
    note_texts = {
        cell.value
        for row in sheet.iter_rows(min_row=totals_row + 1, max_col=1)
        for cell in row
        if cell.value is not None
    }
    assert set(json_body["notes"]) <= note_texts


@pytest.mark.parametrize(
    "name",
    [
        '=HYPERLINK("https://example.test","Amy")',
        "=1+1",
        "+1 555 0100",
        "-Amy-",
        "@amy",
        "#N/A",
    ],
)
def test_a_customer_s_name_is_a_text_cell_whatever_it_begins_with(
    client: TestClient,
    admin_headers: dict[str, str],
    db: Session,
    make_item: ItemFactory,
    admin_user: User,
    name: str,
) -> None:
    """A name a customer chose is text in the workbook, never a formula.

    Read back without evaluating anything: the cell's own type says whether
    a spreadsheet would compute it (`f`), show it as an error (`e`) or show
    it as written (`s`).
    """
    _placed(
        db,
        priced_item(make_item, "Coin", Decimal("40.00")),
        _venue(db, "ebay"),
        admin_user,
        price=Decimal("133.75"),
        status_code="paid",
        buyer_name=name,
    )
    db.commit()

    res = client.get("/api/reports/sl_fulfilment/workbook", headers=admin_headers)
    assert res.status_code == 200, res.text

    sheet = load_workbook(BytesIO(res.content)).worksheets[0]
    named = [cell for row in sheet.iter_rows() for cell in row if cell.value == name]
    assert len(named) == 1
    assert named[0].data_type == "s"


def test_every_text_cell_of_a_report_workbook_is_text() -> None:
    """The title, a parameter, a header, a row, the totals and a note alike."""
    formula = "=1+1"
    report = PR_OUTSTANDING
    result = ReportResult(
        columns=[Column("what", formula, "text"), Column("cost", "Cost", "money")],
        rows=[{"what": formula, "cost": Decimal("1.00")}],
        totals={"what": formula, "cost": Decimal("1.00")},
        drills=[None],
        notes=[formula],
    )

    book = write_workbook(
        report, OutstandingParams(), result, datetime.now().astimezone()
    )

    sheet = book.worksheets[0]
    cells = [cell for row in sheet.iter_rows() for cell in row]
    written = [cell for cell in cells if cell.value == formula]
    # The header, the row, the totals row and the note.
    assert len(written) == 4
    assert {cell.data_type for cell in cells if isinstance(cell.value, str)} == {"s"}
    # A number is still a number, in its own format.
    money = [cell for cell in cells if cell.value == Decimal("1.00")]
    assert [cell.data_type for cell in money] == ["n", "n"]
    assert {cell.number_format for cell in money} == {"#,##0.00"}


def test_the_workbook_formats_percent_and_count_cells(
    client: TestClient, admin_headers: dict[str, str], db: Session
) -> None:
    """`dq_completeness`: a percent column prints "0.0"; a count is still numeric."""
    build_bare_item(db)

    json_res = client.get("/api/reports/dq_completeness", headers=admin_headers)
    assert json_res.status_code == 200
    columns = json_res.json()["columns"]

    wb_res = client.get("/api/reports/dq_completeness/workbook", headers=admin_headers)
    assert wb_res.status_code == 200
    book = load_workbook(BytesIO(wb_res.content))
    sheet = book["Field completeness"]
    header_row = _header_row(sheet, columns[0]["label"])
    data_row = header_row + 1

    count_index = next(
        i for i, c in enumerate(columns, start=1) if c["kind"] == "count"
    )
    percent_index = next(
        i for i, c in enumerate(columns, start=1) if c["kind"] == "percent"
    )

    count_cell = sheet.cell(row=data_row, column=count_index)
    # Written as a float -- openpyxl's own Cell.value has no plain-int
    # member -- but a round trip through the actual .xlsx bytes reads a
    # whole number back as `int`, so only the value, not the Python type,
    # is asserted here.
    assert count_cell.value == 1

    percent_cell = sheet.cell(row=data_row, column=percent_index)
    assert percent_cell.number_format == "0.0"
    assert isinstance(percent_cell.value, (int, float))


def test_the_workbook_lists_its_parameters_then_a_run_at_row(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get("/api/reports/pr_outstanding/workbook", headers=admin_headers)
    assert res.status_code == 200
    book = load_workbook(BytesIO(res.content))
    sheet = book["Not yet arrived"]

    param_row = _header_row(sheet, "Overdue after (days)")
    assert sheet.cell(row=param_row, column=2).value == 21.0

    run_at_row = _header_row(sheet, "Run at")
    assert run_at_row == param_row + 1
    run_at_text = sheet.cell(row=run_at_row, column=2).value
    assert isinstance(run_at_text, str)
    datetime.fromisoformat(run_at_text)  # parses without raising: valid ISO text


def test_the_workbook_freezes_panes_below_the_header(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get("/api/reports/pr_outstanding/workbook", headers=admin_headers)
    assert res.status_code == 200
    book = load_workbook(BytesIO(res.content))
    sheet = book["Not yet arrived"]
    header_row = _header_row(sheet, "Order")
    assert sheet.freeze_panes == f"A{header_row + 1}"


def test_the_workbook_filenames_date_equals_run_ats_date(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get("/api/reports/pr_outstanding/workbook", headers=admin_headers)
    assert res.status_code == 200
    disposition = res.headers["content-disposition"]
    filename_date = disposition.split("pr_outstanding_")[1].removesuffix('.xlsx"')

    book = load_workbook(BytesIO(res.content))
    sheet = book["Not yet arrived"]
    run_at_row = _header_row(sheet, "Run at")
    run_at_text = sheet.cell(row=run_at_row, column=2).value
    assert isinstance(run_at_text, str)
    run_at = datetime.fromisoformat(run_at_text)

    assert filename_date == run_at.date().isoformat()


@pytest.mark.parametrize("report_id", list(REPORTS))
def test_every_registered_report_exports_a_readable_workbook(
    report_id: str, client: TestClient, admin_headers: dict[str, str], db: Session
) -> None:
    """Every report, with its default parameters, exports a workbook that opens.

    Over a collection with rows for the reports to write: the workbook
    refuses a cell whose value does not fit its column's kind, which an
    empty table never asks it to check.
    """
    _build_modest_collection(db)

    json_res = client.get(f"/api/reports/{report_id}", headers=admin_headers)
    assert json_res.status_code == 200, (report_id, json_res.text)
    body = json_res.json()

    res = client.get(f"/api/reports/{report_id}/workbook", headers=admin_headers)
    assert res.status_code == 200, (report_id, res.text)
    book = load_workbook(BytesIO(res.content))
    sheet = book.worksheets[0]
    assert sheet.cell(row=1, column=1).value == REPORTS[report_id].title
    # The table's header sits a blank row below "Run at" -- found from there,
    # since a parameter may carry the same label as the first column.
    header_row = _header_row(sheet, "Run at") + 2
    assert sheet.cell(row=header_row, column=1).value == body["columns"][0]["label"]

    # One sheet row per result row, directly below the header.
    for offset, json_row in enumerate(body["rows"], start=1):
        first = json_row[body["columns"][0]["key"]]
        cell = sheet.cell(row=header_row + offset, column=1).value
        if body["columns"][0]["kind"] == "text":
            assert cell == (first if first != "" else None), (report_id, offset)


def test_the_collection_the_workbooks_are_exported_over_gives_most_reports_rows(
    client: TestClient, admin_headers: dict[str, str], db: Session
) -> None:
    """The export test above means little if the reports it runs are empty."""
    _build_modest_collection(db)

    empty = set()
    for report_id in REPORTS:
        res = client.get(f"/api/reports/{report_id}", headers=admin_headers)
        if not res.json()["rows"]:
            empty.add(report_id)
    # Nothing in that collection is dated outside its series, left for a
    # person to classify, or filled by a rule.
    assert empty <= {"dq_series_years", "dq_series_review", "dq_derived"}


def test_a_date_range_reports_bounds_are_date_cells_and_an_absent_one_is_any(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    """`pr_spend`'s From/To: a date cell when set, the word "any" when not."""
    res = client.get("/api/reports/pr_spend/workbook", headers=admin_headers)
    assert res.status_code == 200, res.text
    sheet = load_workbook(BytesIO(res.content)).worksheets[0]
    assert sheet.cell(row=_header_row(sheet, "From"), column=2).value == "any"
    assert sheet.cell(row=_header_row(sheet, "To"), column=2).value == "any"

    res = client.get(
        "/api/reports/pr_spend/workbook?date_from=2026-01-01&date_to=2026-06-30",
        headers=admin_headers,
    )
    assert res.status_code == 200, res.text
    sheet = load_workbook(BytesIO(res.content)).worksheets[0]
    from_cell = sheet.cell(row=_header_row(sheet, "From"), column=2)
    to_cell = sheet.cell(row=_header_row(sheet, "To"), column=2)
    assert isinstance(from_cell.value, datetime)
    assert from_cell.value.date() == date(2026, 1, 1)
    assert from_cell.number_format == "yyyy-mm-dd"
    assert isinstance(to_cell.value, datetime)
    assert to_cell.value.date() == date(2026, 6, 30)
    assert to_cell.number_format == "yyyy-mm-dd"


def test_a_bad_parameter_or_unknown_report_is_refused_on_the_workbook_path_too(
    client: TestClient, admin_headers: dict[str, str]
) -> None:
    res = client.get(
        "/api/reports/pr_outstanding/workbook?overdue_days=0", headers=admin_headers
    )
    assert res.status_code == 422

    res = client.get(
        "/api/reports/pr_outstanding/workbook?nope=1", headers=admin_headers
    )
    assert res.status_code == 422

    res = client.get("/api/reports/nope/workbook", headers=admin_headers)
    assert res.status_code == 404
