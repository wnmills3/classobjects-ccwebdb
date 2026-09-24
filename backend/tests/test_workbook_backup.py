"""The database as a workbook, and back (`app.workbook_backup`).

A backup the owner can open and correct, then rebuild a database from. The
round trip runs on two scratch databases of its own, with a schema made to
hold every hard case at once: a link to its own table pointing at a row
loaded after it, a link to another table, a computed column, an enum,
JSON with both SQL NULL and JSON's null, a decimal, a zone-aware
timestamp, a date, an empty string beside a NULL, and a value that looks
like a formula.
"""

from __future__ import annotations

from collections.abc import Iterator
from datetime import UTC, date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from app import backup
from app import workbook_backup as wb
from openpyxl import Workbook, load_workbook
from sqlalchemy import (
    JSON,
    Boolean,
    Date,
    DateTime,
    Integer,
    Numeric,
    String,
    create_engine,
    text,
)
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.engine import Engine
from sqlalchemy.types import Enum, TypeEngine

from tests.conftest import TEST_URL

SCHEMA = """
CREATE TYPE shade AS ENUM ('red', 'blue');
CREATE TABLE alembic_version (version_num varchar(32) PRIMARY KEY);
CREATE TABLE parent (
    id serial PRIMARY KEY,
    name text NOT NULL,
    parent_id integer REFERENCES parent(id),
    amount numeric(12, 2),
    doubled numeric(12, 2) GENERATED ALWAYS AS (amount * 2) STORED,
    at timestamptz,
    day date,
    flag boolean,
    data jsonb,
    colour shade
);
CREATE TABLE child (
    id serial PRIMARY KEY,
    parent_id integer NOT NULL REFERENCES parent(id),
    note text
);
"""

ROWS = """
INSERT INTO alembic_version VALUES ('rev_1');
INSERT INTO parent (id, name, parent_id, amount, at, day, flag, data, colour) VALUES
  (1, 'root', NULL, 12.50, '2026-09-24 10:00:00+02', '2026-09-24', true,
   '{"a": [1, 2]}', 'red'),
  -- A child row pointing at a parent with a HIGHER id: loaded in one pass,
  -- its link would reference a row not yet there.
  (2, '', 3, NULL, NULL, NULL, false, NULL, NULL),
  (3, '=SUM(A1:A2)', 1, 0.10, '2026-01-01 00:00:00+00', NULL, NULL, '[]', 'blue'),
  -- JSON's own null, distinct from row 2's SQL NULL.
  (4, 'json null', 1, NULL, NULL, NULL, NULL, 'null', NULL);
SELECT setval('parent_id_seq', 4);
INSERT INTO child (parent_id, note) VALUES (1, NULL), (3, ''), (2, 'x');
"""


def _url(name: str) -> str:
    return TEST_URL.set(database=name).render_as_string(hide_password=False)


@pytest.fixture
def pair(monkeypatch: pytest.MonkeyPatch) -> Iterator[tuple[Engine, str]]:
    """A filled source database and an empty target with the same schema."""
    admin = create_engine(
        TEST_URL.set(database="postgres"), isolation_level="AUTOCOMMIT"
    )
    names = ("ccwebdb_test_wb_source", "ccwebdb_test_wb_target")
    with admin.connect() as conn:
        for name in names:
            conn.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
            conn.execute(text(f'CREATE DATABASE "{name}"'))
    source = create_engine(_url(names[0]))
    target = create_engine(_url(names[1]))
    with source.begin() as conn:
        conn.execute(text(SCHEMA))
        conn.execute(text(ROWS))
    with target.begin() as conn:
        conn.execute(text(SCHEMA))
        conn.execute(text("INSERT INTO alembic_version VALUES ('rev_1')"))
        # As a migration might: a row the import must clear first.
        conn.execute(text("INSERT INTO parent (name) VALUES ('seeded by migration')"))
    # The live database is neither of these.
    monkeypatch.setattr(wb.settings, "database_url", _url("ccwebdb_test_live_stand_in"))
    try:
        yield source, _url(names[1])
    finally:
        source.dispose()
        target.dispose()
        with admin.connect() as conn:
            for name in names:
                conn.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
        admin.dispose()


def test_a_round_trip_rebuilds_the_database_exactly(
    pair: tuple[Engine, str], tmp_path: Path
) -> None:
    source, target_url = pair
    path = tmp_path / "backup.xlsx"
    counts = wb.export_workbook(source, path)
    assert counts == {"parent": 4, "child": 3}

    loaded = wb.import_workbook(path, target_url)
    assert loaded == counts
    target = create_engine(target_url)
    try:
        assert wb.compare(source, target) == []
        with target.connect() as conn:
            # The empty string stayed empty, not NULL; the formula stayed text.
            assert (
                conn.execute(text("SELECT name FROM parent WHERE id = 2")).scalar()
                == ""
            )
            assert (
                conn.execute(text("SELECT note FROM child WHERE id = 2")).scalar() == ""
            )
            assert (
                conn.execute(text("SELECT note FROM child WHERE id = 1")).scalar()
                is None
            )
            # The migration's own row was cleared, and the computed column recomputed.
            assert conn.execute(text("SELECT count(*) FROM parent")).scalar() == 4
            # SQL NULL and JSON's null both survived as what they were.
            nulls = conn.execute(
                text("SELECT data IS NULL FROM parent WHERE id IN (2, 4) ORDER BY id")
            ).scalars()
            assert list(nulls) == [True, False]
            assert conn.execute(
                text("SELECT doubled FROM parent WHERE id = 1")
            ).scalar() == Decimal("25.00")
            # The sequence is past the loaded ids: a new row does not collide.
            new_id = conn.execute(
                text("INSERT INTO parent (name) VALUES ('new') RETURNING id")
            ).scalar()
            assert new_id == 5
    finally:
        target.dispose()


def test_the_workbook_is_readable_and_marks_what_it_cannot_hold(
    pair: tuple[Engine, str], tmp_path: Path
) -> None:
    source, _ = pair
    path = tmp_path / "backup.xlsx"
    wb.export_workbook(source, path)
    book = load_workbook(path)
    assert book.sheetnames[:2] == [wb.ABOUT, wb.COLUMNS]
    # Foreign-key order: the referenced table's sheet comes first.
    assert book.sheetnames.index("parent") < book.sheetnames.index("child")
    sheet = book["parent"]
    header = [c.value for c in sheet[1]]
    assert "doubled (computed)" in header
    row = dict(zip(header, sheet[3], strict=True))  # id 2
    assert row["name"].value == wb.EMPTY_STRING
    formula = dict(zip(header, sheet[4], strict=True))["name"]
    assert formula.value == "=SUM(A1:A2)" and formula.data_type == "s"
    about = {
        r[0]: r[1] for r in book[wb.ABOUT].iter_rows(values_only=True) if r and r[0]
    }
    assert about["migration revision"] == "rev_1"
    assert about["format"] == wb.FORMAT


def test_an_edited_cell_arrives_and_nothing_else_moves(
    pair: tuple[Engine, str], tmp_path: Path
) -> None:
    source, target_url = pair
    path = tmp_path / "backup.xlsx"
    wb.export_workbook(source, path)
    book = load_workbook(path)
    book["child"]["C2"] = "edited"  # child id 1's note
    book.save(path)
    wb.import_workbook(path, target_url)
    target = create_engine(target_url)
    try:
        assert [d.table for d in wb.compare(source, target)] == ["child"]
        with target.connect() as conn:
            assert (
                conn.execute(text("SELECT note FROM child WHERE id = 1")).scalar()
                == "edited"
            )
    finally:
        target.dispose()


def test_a_broken_link_is_refused_and_nothing_is_loaded(
    pair: tuple[Engine, str], tmp_path: Path
) -> None:
    source, target_url = pair
    path = tmp_path / "backup.xlsx"
    wb.export_workbook(source, path)
    book = load_workbook(path)
    book["child"]["B2"] = 999  # no parent 999
    book.save(path)
    with pytest.raises(wb.WorkbookError, match="child: the database refused a row"):
        wb.import_workbook(path, target_url)
    target = create_engine(target_url)
    try:
        with target.connect() as conn:
            # One transaction: the migration's row is still there, untouched.
            assert conn.execute(text("SELECT count(*) FROM parent")).scalar() == 1
    finally:
        target.dispose()


def test_a_mismatched_revision_or_a_missing_column_is_refused(
    pair: tuple[Engine, str], tmp_path: Path
) -> None:
    source, target_url = pair
    path = tmp_path / "backup.xlsx"
    wb.export_workbook(source, path)

    target = create_engine(target_url)
    with target.begin() as conn:
        conn.execute(text("UPDATE alembic_version SET version_num = 'rev_2'"))
    target.dispose()
    with pytest.raises(wb.WorkbookError, match="revision rev_1, the database at rev_2"):
        wb.import_workbook(path, target_url)

    target = create_engine(target_url)
    with target.begin() as conn:
        conn.execute(text("UPDATE alembic_version SET version_num = 'rev_1'"))
    target.dispose()
    book = load_workbook(path)
    book["child"].delete_cols(3)  # the note column
    book.save(path)
    with pytest.raises(wb.WorkbookError, match="child: missing column note"):
        wb.import_workbook(path, target_url)


def test_the_live_database_is_refused(
    pair: tuple[Engine, str], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source, target_url = pair
    path = tmp_path / "backup.xlsx"
    wb.export_workbook(source, path)
    monkeypatch.setattr(wb.settings, "database_url", target_url)
    with pytest.raises(wb.WorkbookError, match="is the live database"):
        wb.import_workbook(path, target_url)


def test_app_backup_finds_the_tables_no_model_describes(
    pair: tuple[Engine, str],
) -> None:
    """`app.backup` finds every table with no model, the migration table too.

    It used to copy the models' tables only, so a restored copy had no import
    history and no migration revision. The scratch database's tables have no
    models at all, so every one of them must be found.
    """
    source, _ = pair
    found = [table.name for table in backup.unmodelled_tables(source)]
    assert set(found) == {"alembic_version", "parent", "child"}
    assert found.index("parent") < found.index("child")
    # And the copy's table list is the models' followed by these.
    names = [table.name for table in backup.all_tables(source)]
    assert names[-len(found) :] == found


def test_remembered_widths_follow_their_column_by_name(
    pair: tuple[Engine, str], tmp_path: Path
) -> None:
    """Widths are keyed by column name, so they land on the right letter.

    `note` is child's third column; its width must reach column C whatever
    position a width file lists it in, and a computed column is sized too.
    """
    source, _ = pair
    path = tmp_path / "sized.xlsx"
    widths = {"child": {"note": 40.5}, "parent": {"doubled (computed)": 13.0}}
    wb.export_workbook(source, path, widths)
    book = load_workbook(path)
    assert book["child"].column_dimensions["C"].width == 40.5
    header = [c.value for c in book["parent"][1]]
    letter = "ABCDEFGHIJ"[header.index("doubled (computed)")]
    assert book["parent"].column_dimensions[letter].width == 13.0
    assert wb.capture_widths(path) == widths


def test_remembering_replaces_a_sized_sheet_and_keeps_the_rest(
    pair: tuple[Engine, str], tmp_path: Path
) -> None:
    source, _ = pair
    store = tmp_path / "widths.json"
    store.write_text(
        '{"child": {"note": 9.0}, "parent": {"name": 9.0, "amount": 9.0}}',
        encoding="utf-8",
    )
    path = tmp_path / "sized.xlsx"
    wb.export_workbook(source, path, {"parent": {"name": 33.0}})
    merged = wb.remember_widths(path, store)
    # parent was resized in the workbook: its entry is the workbook's now.
    assert merged == {"child": {"note": 9.0}, "parent": {"name": 33.0}}
    assert wb.load_widths(store) == merged


def test_a_width_excel_saved_for_a_run_of_columns_reaches_each(
    tmp_path: Path,
) -> None:
    """Excel may store one width for adjacent columns (min=1, max=3)."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "child"
    sheet.append(["id", "parent_id", "note"])
    run = sheet.column_dimensions["A"]
    run.min, run.max, run.width = 1, 3, 17.5
    path = tmp_path / "run.xlsx"
    book.save(path)
    assert wb.capture_widths(path) == {
        "child": {"id": 17.5, "note": 17.5, "parent_id": 17.5}
    }


def test_no_width_file_means_default_widths(tmp_path: Path) -> None:
    assert wb.load_widths(tmp_path / "absent.json") == {}


# -- one cell -------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "kind", "cell"),
    [
        (None, String(), None),
        ("", String(), wb.EMPTY_STRING),
        (Decimal("0.77344"), Numeric(12, 6), 0.77344),
        ({"a": 1}, JSONB(), '{"a": 1}'),
        (
            datetime(2026, 9, 24, 10, 0, tzinfo=timezone(timedelta(hours=2))),
            DateTime(timezone=True),
            "2026-09-24T10:00:00+02:00",
        ),
    ],
)
def test_a_value_is_written_as_a_cell(
    value: object, kind: TypeEngine[Any], cell: object
) -> None:
    assert wb.to_cell(value, kind) == cell


@pytest.mark.parametrize(
    ("cell", "kind", "value"),
    [
        (None, String(), None),
        ("  ", String(), None),
        (wb.EMPTY_STRING, String(), ""),
        (0.77344, Numeric(12, 6), Decimal("0.773440")),
        (12.5, Numeric(12, 2), Decimal("12.50")),
        (5.0, Integer(), 5),
        ("TRUE", Boolean(), True),
        ("no", Boolean(), False),
        (datetime(2026, 9, 24), Date(), date(2026, 9, 24)),
        (
            "2026-09-24T10:00:00+00:00",
            DateTime(timezone=True),
            datetime(2026, 9, 24, 10, tzinfo=UTC),
        ),
        ('{"a": [1]}', JSONB(), {"a": [1]}),
        # JSON's own null is a value, kept apart from SQL NULL (an empty cell).
        ("null", JSONB(), JSON.NULL),
        ("red", Enum("red", "blue"), "red"),
    ],
)
def test_a_cell_is_read_as_its_column_stores_it(
    cell: object, kind: TypeEngine[Any], value: object
) -> None:
    assert wb.from_cell(cell, kind, "t row 2, c") == value


def test_a_bad_cell_names_where_it_is() -> None:
    with pytest.raises(wb.WorkbookError, match="inventory_item row 7, piece_count"):
        wb.from_cell(2.5, Integer(), "inventory_item row 7, piece_count")
