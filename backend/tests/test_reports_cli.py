"""`python -m app.reports list|run` -- the same registry the API uses.

`main(argv, db=...)` is called directly with the test's own `db` fixture
throughout: the CLI's default (no `db` given) opens the application's own
`SessionLocal`, bound to whatever database the settings name -- the live
`ccwebdb` on this machine -- and a test must never do that.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from app.models import ItemStatus
from app.reports import REPORTS
from app.reports.__main__ import main
from app.reports.purchasing import PR_OUTSTANDING, OutstandingParams
from app.reports.serialize import serialize_result
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from tests.builders import build_bare_item, build_purchase_order, code_id


def _outstanding_purchase(db: Session, vendor: str, order_number: str) -> None:
    """A purchase with one live, still-`ordered` item -- what `pr_outstanding` lists."""
    order = build_purchase_order(db, vendor_name=vendor, order_number=order_number)
    item = build_bare_item(
        db,
        status_id=code_id(db, ItemStatus, "ordered"),
        item_cost=Decimal("42.50"),
        shipping_cost=Decimal("0.00"),
        tax_rate=Decimal("0"),
    )
    item.purchase_order_id = order.id
    db.commit()


def test_list_prints_every_registered_id(
    db: Session, capsys: pytest.CaptureFixture[str]
) -> None:
    exit_code = main(["list"], db=db)
    out, err = capsys.readouterr()

    assert exit_code == 0
    assert err == ""
    for report_id in REPORTS:
        assert report_id in out


def test_run_prints_the_table_for_a_built_purchase(
    db: Session, capsys: pytest.CaptureFixture[str]
) -> None:
    _outstanding_purchase(db, "CLI Vendor", "CLI-1")

    exit_code = main(["run", "pr_outstanding"], db=db)
    out, err = capsys.readouterr()

    assert exit_code == 0
    assert err == ""
    assert "Not yet arrived" in out  # the report's title
    assert "CLI-1" in out
    assert "CLI Vendor" in out
    assert "42.50" in out  # money, right-aligned as the Decimal prints


def test_a_name_s_control_characters_are_printed_escaped_on_its_own_line(
    db: Session, capsys: pytest.CaptureFixture[str]
) -> None:
    """A stored name is printed, never obeyed.

    A name may hold an escape sequence that clears the screen or a line
    break that starts a row of its own making. Each such character prints
    as its escape, so the row stays one line and the terminal is sent
    nothing but text.
    """
    _outstanding_purchase(db, "Escape\x1b[2J\r\nVendor‮", "ESC-1")

    exit_code = main(["run", "pr_outstanding"], db=db)
    out, _err = capsys.readouterr()

    assert exit_code == 0
    escaped = "Escape\\x1b[2J\\r\\nVendor\\u202e"
    (line,) = [line for line in out.splitlines() if "ESC-1" in line]
    assert escaped in line
    for character in ("\x1b", "\r", "‮"):
        assert character not in out
    # The column is as wide as what is printed, so the next one still starts
    # under its own header.
    header = next(line for line in out.splitlines() if line.startswith("Order "))
    assert header.index("Seller") == line.index(escaped) + len(escaped) + 2


def test_a_param_is_passed_through(
    db: Session, capsys: pytest.CaptureFixture[str]
) -> None:
    _outstanding_purchase(db, "Param Vendor", "PARAM-1")

    exit_code = main(["run", "pr_outstanding", "--param", "overdue_days=5"], db=db)
    out, _err = capsys.readouterr()

    assert exit_code == 0
    assert "Overdue after (days): 5" in out


def test_a_bad_parameter_value_exits_2(
    db: Session, capsys: pytest.CaptureFixture[str]
) -> None:
    exit_code = main(["run", "pr_outstanding", "--param", "overdue_days=0"], db=db)
    out, err = capsys.readouterr()

    assert exit_code == 2
    assert out == ""
    assert err != ""


def test_an_unknown_parameter_name_exits_2(
    db: Session, capsys: pytest.CaptureFixture[str]
) -> None:
    exit_code = main(["run", "pr_outstanding", "--param", "bogus=1"], db=db)
    _out, err = capsys.readouterr()

    assert exit_code == 2
    assert "bogus" in err


def test_a_malformed_param_with_no_equals_sign_exits_2(
    db: Session, capsys: pytest.CaptureFixture[str]
) -> None:
    exit_code = main(["run", "pr_outstanding", "--param", "overdue_days"], db=db)
    _out, err = capsys.readouterr()

    assert exit_code == 2
    assert err != ""


def test_an_unknown_report_id_exits_2_and_lists_valid_ids(
    db: Session, capsys: pytest.CaptureFixture[str]
) -> None:
    exit_code = main(["run", "nope"], db=db)
    out, err = capsys.readouterr()

    assert exit_code == 2
    assert out == ""
    assert "pr_outstanding" in err


def _header_row(sheet: Worksheet, first_label: str) -> int:
    """The row number whose first cell is `first_label` -- the header row."""
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=1):
        if row[0].value == first_label:
            return row[0].row
    raise AssertionError(f"no row starts with {first_label!r}")


def test_workbook_data_rows_equal_the_serialized_rows(
    db: Session, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    _outstanding_purchase(db, "WB Vendor", "WB-CLI-1")
    path = tmp_path / "report.xlsx"

    exit_code = main(["run", "pr_outstanding", "--workbook", str(path)], db=db)
    out, err = capsys.readouterr()

    assert exit_code == 0
    assert err == ""
    assert str(path) in out
    assert path.exists()

    # The same run, replayed for its JSON shape -- the same comparison
    # test_reports_api.py's workbook test makes against the HTTP endpoint.
    result = PR_OUTSTANDING.run(db, OutstandingParams())
    body = serialize_result(
        PR_OUTSTANDING, OutstandingParams(), result, datetime.now().astimezone()
    )
    columns = body["columns"]
    assert isinstance(columns, list)
    rows = body["rows"]
    assert isinstance(rows, list)

    book = load_workbook(path)
    sheet = book.active
    assert sheet is not None
    header_row = _header_row(sheet, columns[0]["label"])

    for offset, json_row in enumerate(rows, start=1):
        for col_index, column in enumerate(columns, start=1):
            cell = sheet.cell(row=header_row + offset, column=col_index)
            json_value = json_row[column["key"]]
            if column["kind"] == "money":
                assert Decimal(str(cell.value)) == Decimal(json_value)
            elif column["kind"] == "date":
                if json_value is None:
                    assert cell.value is None
                else:
                    assert isinstance(cell.value, (date, datetime))
                    assert cell.value.isoformat() == json_value
            else:
                expected = json_value if json_value != "" else None
                assert cell.value == expected
